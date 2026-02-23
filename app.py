import re
import io
import pandas as pd
import fitz  # pymupdf
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title="Extrator de Pedidos (Atacadão)", layout="wide")

def norm_line(ln: str) -> str:
    ln = ln.strip()
    if ln.startswith("|"):
        ln = ln.strip("|").strip()
    ln = re.sub(r"\s+", " ", ln)
    return ln

def clean_desc(desc: str) -> str:
    desc = re.sub(r"^RF[.\s-]*", "", desc.strip(), flags=re.IGNORECASE)
    desc = re.sub(r"\s{2,}", " ", desc).strip()
    return desc

cnpj_re = re.compile(r"Local de Entrega:\s*([0-9]{8}/[0-9]{4}-[0-9]{2})")
pedido_re = re.compile(r"Numero:\s*([0-9]+)")

item_line_re = re.compile(
    r"^(?P<desc>.+?)\s+(?P<dt>\d{2}/\d{2})\s+(?P<qtde>\d+)\s+(?P<unit>\d+,\d+)\s+.*\s+(?P<peso>\d+,\d+)\s+[A-Z]$"
)

code_uom_pack_re = re.compile(r"(?P<cod>\d{8}/\d{3})\s+(?P<uom>CXA|KG)\s+(?P<pack>\S+)")

def parse_pdf_bytes(pdf_bytes: bytes) -> list[dict]:
    rows = []
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")

    for page in doc:
        raw_lines = [ln for ln in page.get_text("text").splitlines() if ln.strip()]
        lines = [norm_line(ln) for ln in raw_lines if not ln.strip().startswith("+")]

        cnpj = None
        pedido = None
        for ln in lines:
            if cnpj is None:
                m = cnpj_re.search(ln)
                if m:
                    cnpj = m.group(1)
            if pedido is None:
                m = pedido_re.search(ln)
                if m:
                    pedido = m.group(1)
            if cnpj and pedido:
                break

        i = 0
        while i < len(lines):
            m = item_line_re.match(lines[i])
            if m:
                qtde = int(m.group("qtde"))
                if qtde == 999:
                    i += 1
                    continue

                desc = clean_desc(m.group("desc"))
                valor = float(m.group("unit").replace(".", "").replace(",", "."))
                peso_total = float(m.group("peso").replace(".", "").replace(",", "."))

                cod = ""
                uom = ""
                pack = ""
                if i + 1 < len(lines):
                    mc = code_uom_pack_re.search(lines[i + 1])
                    if mc:
                        cod = mc.group("cod")
                        uom = mc.group("uom")
                        pack = mc.group("pack")

                rows.append({
                    "CNPJ": cnpj,
                    "Pedido": pedido,
                    "Código do produto": cod,
                    "Descrição do Produto": desc,
                    "Peso Unitário": pack,
                    "Unidade de Medida": uom,
                    "Quantidade": qtde,
                    "Valor": valor,
                    "Peso": peso_total,
                })
            i += 1

    return rows

def build_excel(df: pd.DataFrame) -> bytes:
    output = io.BytesIO()

    # Exporta com pandas
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Itens")

    # Ajusta formatação e cria Valor total com openpyxl
    output.seek(0)
    wb = load_workbook(output)
    ws = wb["Itens"]

    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
    col_qtde = headers["Quantidade"]
    col_valor = headers["Valor"]

    ws.insert_cols(col_valor + 1)
    ws.cell(row=1, column=col_valor + 1).value = "Valor total"

    for r in range(2, ws.max_row + 1):
        qtde_cell = ws.cell(row=r, column=col_qtde)
        valor_cell = ws.cell(row=r, column=col_valor)
        total_cell = ws.cell(row=r, column=col_valor + 1)

        total_cell.value = f"={qtde_cell.coordinate}*{valor_cell.coordinate}"
        valor_cell.number_format = 'R$ #,##0.00'
        total_cell.number_format = 'R$ #,##0.00'

    out2 = io.BytesIO()
    wb.save(out2)
    return out2.getvalue()

# =========================
# UI (layout e textos)
# =========================

logo_url = "https://logodownload.org/wp-content/uploads/2018/06/atacadao-logo.png"

# Logo no canto superior esquerdo + Título ao lado
col_logo, col_title = st.columns([1, 7], vertical_alignment="center")
with col_logo:
    st.image(logo_url, width=140)
with col_title:
    st.title("Extrator de Pedidos (Atacadão)")
    st.markdown("<small><i>Criado por Caiã Ricardo Grade.</i></small>", unsafe_allow_html=True)

# Passo a passo quebrado em linhas
st.markdown(
    """
1) Faça upload dos PDFs  
2) Clique em **Extrair pedidos**  
3) Baixe o Excel.
"""
)

pdfs = st.file_uploader("Upload de PDFs", type=["pdf"], accept_multiple_files=True)

run = st.button("🚀 Extrair pedidos", use_container_width=True, disabled=not pdfs)

if run:
    all_rows = []
    with st.spinner("Processando PDFs..."):
        for f in pdfs:
            all_rows.extend(parse_pdf_bytes(f.getvalue()))

        df = pd.DataFrame(all_rows, columns=[
            "CNPJ","Pedido","Código do produto","Descrição do Produto",
            "Peso Unitário","Unidade de Medida",
            "Quantidade","Valor","Peso"
        ])

        if df.empty:
            st.error("Não consegui extrair itens desses PDFs. Eles podem estar escaneados (imagem) ou fora do padrão.")
        else:
            excel_bytes = build_excel(df)
            st.success(f"✅ Concluído: {len(df)} itens extraídos.")
            st.dataframe(df, use_container_width=True, height=320)

            st.download_button(
                label="⬇️ Baixar Excel",
                data=excel_bytes,
                file_name="pedidos_extraidos.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

            st.success("Conversão realizada!!!")
